import csv
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO, StringIO
from typing import Annotated, cast, final, override

from grafy_core.artifacts import ArtifactRef, NodeConfig, NodeInput, NodeOutput
from grafy_core.file_artifacts import load_file_artifact
from grafy_core.file_contracts import CSV_FILE, XLSX_FILE
from grafy_core.nodes import InPort, Node, NodeExecutionContext, OutPort
from grafy_core.plugins import NodeCachePolicy
from grafy_core.ports.artifacts import UnitOfWorkPort
from grafy_core.ports.storage import FileStoragePort
from grafy_core.table_contracts import (
    TABLE_DATA,
    Table,
    TableColumn,
    TableValue,
    TableValueType,
)
from openpyxl import load_workbook
from pydantic import (
    Field,
    StrictBool,
    StrictInt,
    StrictStr,
    field_validator,
)

from grafy_plugin_table.declaration import TABLES


class TableFileImportError(RuntimeError):
    pass


def _table_value(value: object) -> TableValue:
    if value is None or isinstance(value, str | bool):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return value
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, datetime | date):
        return value.isoformat()
    return str(value)


def _inferred_table_value_type(values: list[TableValue]) -> TableValueType:
    observed: set[TableValueType] = set()
    for value in values:
        if value is None:
            continue
        if isinstance(value, bool):
            observed.add(TableValueType.BOOLEAN)
        elif isinstance(value, int):
            observed.add(TableValueType.INTEGER)
        elif isinstance(value, float):
            observed.add(TableValueType.NUMBER)
        elif isinstance(value, str):
            observed.add(TableValueType.TEXT)
        else:
            observed.add(TableValueType.JSON)
    if not observed:
        return TableValueType.UNKNOWN
    if observed <= {TableValueType.INTEGER, TableValueType.NUMBER}:
        return (
            TableValueType.INTEGER
            if observed == {TableValueType.INTEGER}
            else TableValueType.NUMBER
        )
    return next(iter(observed)) if len(observed) == 1 else TableValueType.MIXED


def _table_from_matrix(
    matrix: list[list[object]],
    *,
    header_row: int,
    skip_empty_rows: bool,
) -> Table:
    header_index = header_row - 1
    if header_index >= len(matrix):
        raise TableFileImportError(
            f"Header row {header_row} is outside the imported file"
        )
    data_rows = matrix[header_index + 1 :]
    column_count = max(
        [len(matrix[header_index]), *(len(row) for row in data_rows)],
        default=0,
    )
    if column_count == 0:
        raise TableFileImportError("The imported file does not contain any columns")

    header = matrix[header_index]
    columns = [
        TableColumn(
            id=f"column_{column_index + 1}",
            title=(
                str(header[column_index]).strip()
                if column_index < len(header)
                and header[column_index] is not None
                and str(header[column_index]).strip()
                else f"Column {column_index + 1}"
            ),
        )
        for column_index in range(column_count)
    ]
    rows: list[dict[str, TableValue]] = []
    for raw_row in data_rows:
        values = [
            _table_value(raw_row[column_index] if column_index < len(raw_row) else None)
            for column_index in range(column_count)
        ]
        if skip_empty_rows and all(value is None or value == "" for value in values):
            continue
        rows.append(
            {column.id: value for column, value in zip(columns, values, strict=True)}
        )

    typed_columns = [
        column.model_copy(
            update={
                "value_type": _inferred_table_value_type(
                    [row[column.id] for row in rows]
                )
            }
        )
        for column in columns
    ]
    return Table(columns=typed_columns, rows=rows)


def _csv_matrix(content: bytes, delimiter: str | None) -> list[list[object]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise TableFileImportError("CSV files must use UTF-8 encoding") from exc
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:16_384]).delimiter
        except csv.Error:
            delimiter = ","
    return [
        cast(list[object], row)
        for row in csv.reader(StringIO(text), delimiter=delimiter)
    ]


def _xlsx_matrix(
    content: bytes,
    *,
    sheet_name: str | None,
) -> list[list[object]]:
    try:
        workbook = load_workbook(
            filename=BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise TableFileImportError("The file is not a readable XLSX workbook") from exc
    try:
        if sheet_name is not None:
            if sheet_name not in workbook.sheetnames:
                raise TableFileImportError(
                    f"Workbook has no worksheet named {sheet_name!r}"
                )
            worksheet = workbook[sheet_name]
        else:
            worksheet = workbook.active
        if worksheet is None:
            raise TableFileImportError("The workbook does not have an active worksheet")
        return [
            [cast(object, value) for value in row]
            for row in worksheet.iter_rows(values_only=True)
        ]
    finally:
        workbook.close()


class TableImportConfig(NodeConfig):
    delimiter: StrictStr | None = Field(
        default=None,
        min_length=1,
        max_length=1,
        description="CSV delimiter. Leave empty to detect it from the file.",
    )
    header_row: StrictInt = Field(
        default=1,
        ge=1,
        description="One-based row containing column titles.",
    )
    sheet_name: StrictStr | None = Field(
        default=None,
        min_length=1,
        max_length=255,
        description="XLSX worksheet name. Leave empty to use the active sheet.",
    )
    skip_empty_rows: StrictBool = True

    @field_validator("sheet_name")
    @classmethod
    def validate_sheet_name(cls, value: str | None) -> str | None:
        if value is not None and value != value.strip():
            raise ValueError("sheet_name must not have surrounding whitespace")
        return value


class TableImportInput(NodeInput):
    file: Annotated[
        ArtifactRef,
        InPort(CSV_FILE, also_accepts=(XLSX_FILE,)),
        Field(description="CSV or XLSX file artifact to import."),
    ]


class TableImportOutput(NodeOutput):
    table: Annotated[
        Table,
        OutPort(TABLE_DATA),
        Field(description="Table imported from the selected worksheet or CSV file."),
    ]


@TABLES.node(
    operator_id="table.import",
    version=1,
    title="Import table",
    factory=lambda context: ImportTableNode(
        storage=context.storage,
        uow=context.uow,
    ),
    cache_policy=NodeCachePolicy.NEVER,
)
@final
class ImportTableNode(Node[TableImportConfig, TableImportInput, TableImportOutput]):
    """Import one persisted CSV or XLSX file artifact as a table."""

    def __init__(self, *, storage: FileStoragePort, uow: UnitOfWorkPort) -> None:
        self._storage = storage
        self._uow = uow

    @override
    async def run(
        self,
        context: NodeExecutionContext,
        config: TableImportConfig,
        inputs: TableImportInput,
        /,
    ) -> TableImportOutput:
        ref = inputs.file
        file = await load_file_artifact(
            storage=self._storage,
            uow=self._uow,
            workspace_id=context.workspace_id,
            ref=ref,
        )
        if ref.artifact_type == CSV_FILE.key.id:
            matrix = _csv_matrix(file.content, config.delimiter)
        elif ref.artifact_type == XLSX_FILE.key.id:
            matrix = _xlsx_matrix(file.content, sheet_name=config.sheet_name)
        else:
            raise TableFileImportError(
                f"Table import does not accept {ref.artifact_type}@"
                f"{ref.schema_version} for artifact {ref.artifact_id}"
            )
        return TableImportOutput(
            table=_table_from_matrix(
                matrix,
                header_row=config.header_row,
                skip_empty_rows=config.skip_empty_rows,
            )
        )
