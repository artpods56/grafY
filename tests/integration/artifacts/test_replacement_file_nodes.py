"""The four replacement file nodes turn a persisted file.* artifact into a value.

Ingest writes the bytes and never interprets them, so each test seeds the exact
`file.*` artifact the upload path produces and then runs the visible node.
"""

import json
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from uuid import UUID

import pytest
from grafy_core.artifact_contracts import RASTER_IMAGE, RasterImageContent
from grafy_core.artifacts import (
    ArtifactObject,
    ArtifactRef,
    ArtifactRefSequence,
    ArtifactTypeSpec,
    JsonObject,
)
from grafy_core.domain.plugin_capabilities import PluginRuntimeCapability
from grafy_core.file_artifacts import FileArtifactError
from grafy_core.file_contracts import (
    BMP_FILE,
    CSV_FILE,
    GEOJSON_FILE,
    JPEG_FILE,
    JSON_FILE,
    PNG_FILE,
    TIFF_FILE,
    WEBP_FILE,
    XLSX_FILE,
)
from grafy_core.nodes import NodeExecutionContext
from grafy_core.plugins import PluginRegistry
from grafy_core.ports.storage import SaveFileCommand
from grafy_core.runtime.execution import NodeRuntime
from grafy_core.runtime.in_memory import InMemoryUnitOfWork
from grafy_core.runtime.materialization import InputMaterializer
from grafy_core.runtime.persistence import (
    ArtifactWriterRegistry,
    OutputPersister,
    PersistedNodeOutput,
)
from grafy_core.runtime.resolvers import ResolverRegistry
from grafy_core.table_contracts import TABLE_DATA
from grafy_plugin_gis.artifacts import (
    GEO_FEATURE_COLLECTION,
    GEO_RASTER_SCAN,
)
from grafy_plugin_gis.nodes import (
    GeoJsonParseError,
    GeoJsonParseInput,
    GeoRasterScanImportConfig,
    GeoRasterScanImportInput,
    ImportRasterScanNode,
    ParseGeoJsonNode,
)
from grafy_plugin_gis.plugin import GIS
from grafy_storage import LocalFileObjectStore
from grafy_workbench.image.declaration import IMAGES
from grafy_workbench.image.nodes import (
    DecodeImagesNode,
    ImageDecodeInput,
    RasterImageOutputWriter,
)
from grafy_workbench.table.declaration import TABLES
from grafy_workbench.table.nodes import (
    ImportTableNode,
    TableFileImportError,
    TableImportConfig,
    TableImportInput,
)
from openpyxl import Workbook

TEST_WORKSPACE_ID = UUID("00000000-0000-4000-8000-000000000930")


def feature_collection_bytes(*coordinates: tuple[float, float]) -> bytes:
    return json.dumps(
        {
            "type": "FeatureCollection",
            "features": [
                {
                    "type": "Feature",
                    "id": index,
                    "properties": {"name": f"place-{index}"},
                    "geometry": {"type": "Point", "coordinates": list(position)},
                }
                for index, position in enumerate(coordinates)
            ],
        }
    ).encode("utf-8")


def workbook_bytes(
    *,
    sheets: dict[str, list[list[object]]],
    active: str,
) -> bytes:
    workbook = Workbook()
    empty = workbook.active
    assert empty is not None
    workbook.remove(empty)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    workbook.active = workbook.sheetnames.index(active)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


async def seed_file_artifact(
    storage: LocalFileObjectStore,
    uow: InMemoryUnitOfWork,
    *,
    spec: ArtifactTypeSpec,
    content: bytes,
    original_filename: str | None = None,
    content_type: str = "application/octet-stream",
) -> ArtifactRef:
    """Persist the artifact and bytes ingest would have written for this format."""

    stored = await storage.save(
        SaveFileCommand(
            bucket="artifacts",
            path=f"files/{spec.key.id}/{sha256(content).hexdigest()}",
            stream=BytesIO(content),
            content_type=content_type,
            metadata={},
            allow_overwrite=True,
        )
    )
    metadata: JsonObject = {}
    if original_filename is not None:
        metadata["original_filename"] = original_filename
    artifact = ArtifactObject(
        workspace_id=TEST_WORKSPACE_ID,
        artifact_type=spec.key.id,
        schema_version=spec.key.schema_version,
        content_type=content_type,
        bucket=stored.bucket,
        object_key=stored.path,
        byte_size=stored.byte_size,
        sha256=stored.sha256,
        metadata=metadata,
    )
    async with uow as entered:
        await entered.artifacts.add(artifact)
        await entered.commit()
    return artifact.ref()


def node_context() -> NodeExecutionContext:
    return NodeExecutionContext(
        workspace_id=TEST_WORKSPACE_ID,
        node_id="replacement-file-node",
    )


@pytest.mark.asyncio
async def test_image_decode_decodes_the_whole_ordered_batch_through_the_runtime(
    tmp_path: Path,
) -> None:
    storage = LocalFileObjectStore(tmp_path / "objects")
    uow = InMemoryUnitOfWork()
    first = await seed_file_artifact(
        storage,
        uow,
        spec=PNG_FILE,
        content=b"\x89PNG\r\n\x1a\nfirst",
        original_filename="page-001.png",
    )
    second = await seed_file_artifact(
        storage,
        uow,
        spec=PNG_FILE,
        content=b"\x89PNG\r\n\x1a\nsecond",
        original_filename="page-002.png",
    )
    runtime = NodeRuntime(
        materializer=InputMaterializer(ResolverRegistry()),
        persister=OutputPersister(
            ArtifactWriterRegistry(
                [RasterImageOutputWriter(storage=storage, uow=uow, bucket="artifacts")]
            )
        ),
    )

    output = await runtime.run_node(
        DecodeImagesNode(storage=storage, uow=uow),
        node_context(),
        {
            "files": ArtifactRefSequence.from_key(
                key=PNG_FILE.key,
                item_refs=[second, first],
            )
        },
    )

    assert isinstance(output, PersistedNodeOutput)
    images = output["images"]
    assert isinstance(images, ArtifactRefSequence)
    assert images.artifact_type == RASTER_IMAGE.key.id
    assert len(images.item_refs) == 2
    async with uow as entered:
        stored_images = await entered.artifacts.list_by_type(
            TEST_WORKSPACE_ID,
            RASTER_IMAGE.key,
        )
    assert [artifact.sha256 for artifact in stored_images] == [
        sha256(b"\x89PNG\r\n\x1a\nsecond").hexdigest(),
        sha256(b"\x89PNG\r\n\x1a\nfirst").hexdigest(),
    ]
    assert [artifact.metadata["original_filename"] for artifact in stored_images] == [
        "page-002.png",
        "page-001.png",
    ]


@pytest.mark.parametrize(
    ("spec", "content_type"),
    [
        (PNG_FILE, "image/png"),
        (JPEG_FILE, "image/jpeg"),
        (WEBP_FILE, "image/webp"),
        (TIFF_FILE, "image/tiff"),
        (BMP_FILE, "image/bmp"),
    ],
)
@pytest.mark.asyncio
async def test_image_decode_maps_each_accepted_container_to_its_content_type(
    tmp_path: Path,
    spec: ArtifactTypeSpec,
    content_type: str,
) -> None:
    storage = LocalFileObjectStore(tmp_path / "objects")
    uow = InMemoryUnitOfWork()
    ref = await seed_file_artifact(storage, uow, spec=spec, content=b"image-bytes")
    node = DecodeImagesNode(storage=storage, uow=uow)

    output = await node.run(
        node_context(),
        node.config_contract.model.model_validate({}),
        ImageDecodeInput(
            files=ArtifactRefSequence.from_key(key=spec.key, item_refs=[ref])
        ),
    )

    assert len(output.images) == 1
    assert isinstance(output.images[0], RasterImageContent)
    assert output.images[0].content_type == content_type
    assert output.images[0].filename is None


@pytest.mark.asyncio
async def test_table_import_dispatches_on_the_declared_format_not_the_filename(
    tmp_path: Path,
) -> None:
    storage = LocalFileObjectStore(tmp_path / "objects")
    uow = InMemoryUnitOfWork()
    csv_ref = await seed_file_artifact(
        storage,
        uow,
        spec=CSV_FILE,
        content=b"name,population\nBelynichi,10\n",
        original_filename="places.xlsx",
    )
    xlsx_ref = await seed_file_artifact(
        storage,
        uow,
        spec=XLSX_FILE,
        content=workbook_bytes(
            sheets={"Places": [["name", "population"], ["Belynichi", 10]]},
            active="Places",
        ),
        original_filename="places.csv",
    )
    node = ImportTableNode(storage=storage, uow=uow)

    csv_output = await node.run(
        node_context(),
        node.config_contract.model.model_validate({}),
        TableImportInput(file=csv_ref),
    )
    xlsx_output = await node.run(
        node_context(),
        node.config_contract.model.model_validate({}),
        TableImportInput(file=xlsx_ref),
    )

    assert [column.title for column in csv_output.table.columns] == [
        "name",
        "population",
    ]
    assert csv_output.table.rows == [
        {"column_1": "Belynichi", "column_2": "10"},
    ]
    assert [column.title for column in xlsx_output.table.columns] == [
        "name",
        "population",
    ]
    assert xlsx_output.table.rows == [
        {"column_1": "Belynichi", "column_2": 10},
    ]


@pytest.mark.asyncio
async def test_table_import_applies_delimiter_header_and_empty_row_parameters(
    tmp_path: Path,
) -> None:
    storage = LocalFileObjectStore(tmp_path / "objects")
    uow = InMemoryUnitOfWork()
    ref = await seed_file_artifact(
        storage,
        uow,
        spec=CSV_FILE,
        content=b"Places of the region\nname;population\nBelynichi;10\n;\n",
    )
    node = ImportTableNode(storage=storage, uow=uow)

    default_output = await node.run(
        node_context(),
        node.config_contract.model.model_validate({}),
        TableImportInput(file=ref),
    )
    configured_output = await node.run(
        node_context(),
        TableImportConfig(delimiter=";", header_row=2, skip_empty_rows=True),
        TableImportInput(file=ref),
    )
    kept_output = await node.run(
        node_context(),
        TableImportConfig(delimiter=";", header_row=2, skip_empty_rows=False),
        TableImportInput(file=ref),
    )

    assert default_output.table.columns[0].title == "Places of the region"
    assert [column.title for column in configured_output.table.columns] == [
        "name",
        "population",
    ]
    assert configured_output.table.rows == [
        {"column_1": "Belynichi", "column_2": "10"},
    ]
    assert kept_output.table.rows == [
        {"column_1": "Belynichi", "column_2": "10"},
        {"column_1": "", "column_2": ""},
    ]


@pytest.mark.asyncio
async def test_table_import_selects_the_named_sheet_and_rejects_a_zip_without_one(
    tmp_path: Path,
) -> None:
    storage = LocalFileObjectStore(tmp_path / "objects")
    uow = InMemoryUnitOfWork()
    ref = await seed_file_artifact(
        storage,
        uow,
        spec=XLSX_FILE,
        content=workbook_bytes(
            sheets={
                "Ignored": [["ignored_header"], ["ignored_row"]],
                "Places": [["name", "population"], ["Belynichi", 10]],
            },
            active="Ignored",
        ),
    )
    broken_ref = await seed_file_artifact(
        storage,
        uow,
        spec=XLSX_FILE,
        content=b"PK\x03\x04not-a-workbook",
    )
    node = ImportTableNode(storage=storage, uow=uow)

    active = await node.run(
        node_context(),
        TableImportConfig(),
        TableImportInput(file=ref),
    )
    selected = await node.run(
        node_context(),
        TableImportConfig(sheet_name="Places"),
        TableImportInput(file=ref),
    )
    assert [column.title for column in active.table.columns] == ["ignored_header"]
    assert active.table.rows == [{"column_1": "ignored_row"}]
    assert selected.table.rows == [{"column_1": "Belynichi", "column_2": 10}]

    with pytest.raises(TableFileImportError, match="no worksheet named 'Missing'"):
        _ = await node.run(
            node_context(),
            TableImportConfig(sheet_name="Missing"),
            TableImportInput(file=ref),
        )
    with pytest.raises(TableFileImportError, match="not a readable XLSX workbook"):
        _ = await node.run(
            node_context(),
            TableImportConfig(),
            TableImportInput(file=broken_ref),
        )


@pytest.mark.asyncio
async def test_geojson_parse_accepts_both_declared_container_types(
    tmp_path: Path,
) -> None:
    storage = LocalFileObjectStore(tmp_path / "objects")
    uow = InMemoryUnitOfWork()
    geojson_ref = await seed_file_artifact(
        storage,
        uow,
        spec=GEOJSON_FILE,
        content=feature_collection_bytes((13.405, 52.52), (2.3522, 48.8566)),
        original_filename="cities.geojson",
    )
    json_ref = await seed_file_artifact(
        storage,
        uow,
        spec=JSON_FILE,
        content=feature_collection_bytes((13.405, 52.52)),
        original_filename="cities.json",
    )
    node = ParseGeoJsonNode(storage=storage, uow=uow)

    geojson = await node.run(
        node_context(),
        node.config_contract.model.model_validate({}),
        GeoJsonParseInput(file=geojson_ref),
    )
    plain_json = await node.run(
        node_context(),
        node.config_contract.model.model_validate({}),
        GeoJsonParseInput(file=json_ref),
    )

    assert geojson.features.source_name == "cities.geojson"
    assert geojson.features.bounds == (2.3522, 48.8566, 13.405, 52.52)
    assert plain_json.features.source_name == "cities.json"
    assert plain_json.features.bounds == (13.405, 52.52, 13.405, 52.52)


@pytest.mark.asyncio
async def test_geojson_parse_validates_the_document_shape_at_run_time(
    tmp_path: Path,
) -> None:
    storage = LocalFileObjectStore(tmp_path / "objects")
    uow = InMemoryUnitOfWork()
    array_ref = await seed_file_artifact(
        storage,
        uow,
        spec=JSON_FILE,
        content=b"[1, 2, 3]",
        original_filename="numbers.json",
    )
    projected_ref = await seed_file_artifact(
        storage,
        uow,
        spec=GEOJSON_FILE,
        content=feature_collection_bytes((4_000_000.0, 5_000_000.0)),
        original_filename="projected.geojson",
    )
    node = ParseGeoJsonNode(storage=storage, uow=uow)

    with pytest.raises(GeoJsonParseError, match="not a valid GeoJSON"):
        _ = await node.run(
            node_context(),
            node.config_contract.model.model_validate({}),
            GeoJsonParseInput(file=array_ref),
        )
    with pytest.raises(GeoJsonParseError, match="outside WGS84"):
        _ = await node.run(
            node_context(),
            node.config_contract.model.model_validate({}),
            GeoJsonParseInput(file=projected_ref),
        )


@pytest.mark.asyncio
async def test_raster_scan_import_names_the_source_and_keeps_the_original_filename(
    tmp_path: Path,
) -> None:
    storage = LocalFileObjectStore(tmp_path / "objects")
    uow = InMemoryUnitOfWork()
    named_ref = await seed_file_artifact(
        storage,
        uow,
        spec=TIFF_FILE,
        content=b"II*\x00geotiff-bytes",
        original_filename="scan.tif",
    )
    unnamed_ref = await seed_file_artifact(
        storage,
        uow,
        spec=TIFF_FILE,
        content=b"II*\x00geotiff-bytes-again",
    )
    node = ImportRasterScanNode(storage=storage, uow=uow)

    named = await node.run(
        node_context(),
        GeoRasterScanImportConfig(source_name="Orthophoto 2024"),
        GeoRasterScanImportInput(file=named_ref),
    )
    unnamed = await node.run(
        node_context(),
        GeoRasterScanImportConfig(source_name="Orthophoto 2024"),
        GeoRasterScanImportInput(file=unnamed_ref),
    )

    assert named.raster.content == b"II*\x00geotiff-bytes"
    assert named.raster.source_name == "Orthophoto 2024"
    assert named.raster.filename == "scan.tif"
    assert unnamed.raster.filename == "Orthophoto 2024"


@pytest.mark.asyncio
async def test_file_artifact_reader_fails_closed_without_stored_bytes() -> None:
    storage = LocalFileObjectStore(Path("/tmp/grafy-missing-store"))
    uow = InMemoryUnitOfWork()
    artifact = ArtifactObject(
        workspace_id=TEST_WORKSPACE_ID,
        artifact_type=CSV_FILE.key.id,
        schema_version=CSV_FILE.key.schema_version,
        content_type="text/csv",
    )
    async with uow as entered:
        await entered.artifacts.add(artifact)
        await entered.commit()
    node = ImportTableNode(storage=storage, uow=uow)

    with pytest.raises(FileArtifactError, match="has no stored object"):
        _ = await node.run(
            node_context(),
            TableImportConfig(),
            TableImportInput(file=artifact.ref()),
        )


def test_replacement_nodes_declare_their_catalog_contracts() -> None:
    registry = PluginRegistry()
    registry.install(IMAGES)
    registry.install(TABLES)
    registry.install(GIS)
    registry.freeze()

    decode = registry.node_registration("image.decode", 1)
    decode_files = decode.node_class.input_contract.ports["files"]
    assert decode_files.accepts == PNG_FILE.key
    assert decode_files.also_accepts == (
        JPEG_FILE.key,
        WEBP_FILE.key,
        TIFF_FILE.key,
        BMP_FILE.key,
    )
    assert decode_files.shape.value == "many"
    assert decode_files.preserves_ref_container is True
    decode_images = decode.node_class.output_contract.ports["images"]
    assert decode_images.produces == RASTER_IMAGE.key
    assert decode_images.shape.value == "many"
    assert decode.staged_upload_inputs == ()
    assert decode.required_capabilities == ()

    import_table = registry.node_registration("table.import", 1)
    table_file = import_table.node_class.input_contract.ports["file"]
    assert table_file.accepts == CSV_FILE.key
    assert table_file.also_accepts == (XLSX_FILE.key,)
    assert table_file.shape.value == "one"
    assert table_file.preserves_ref_container is True
    assert (
        import_table.node_class.output_contract.ports["table"].produces
        == TABLE_DATA.key
    )
    assert list(import_table.node_class.config_contract.model.model_fields) == [
        "delimiter",
        "header_row",
        "sheet_name",
        "skip_empty_rows",
    ]
    assert import_table.staged_upload_inputs == ()
    assert import_table.required_capabilities == ()

    parse_geojson = registry.node_registration("gis.geojson.parse", 1)
    geojson_file = parse_geojson.node_class.input_contract.ports["file"]
    assert geojson_file.accepts == GEOJSON_FILE.key
    assert geojson_file.also_accepts == (JSON_FILE.key,)
    assert geojson_file.shape.value == "one"
    assert (
        parse_geojson.node_class.output_contract.ports["features"].produces
        == GEO_FEATURE_COLLECTION.key
    )
    assert parse_geojson.staged_upload_inputs == ()
    assert parse_geojson.required_capabilities == ()

    import_raster = registry.node_registration("gis.raster_scan.import", 1)
    raster_file = import_raster.node_class.input_contract.ports["file"]
    assert raster_file.accepts == TIFF_FILE.key
    assert raster_file.also_accepts == ()
    assert raster_file.shape.value == "one"
    assert (
        import_raster.node_class.output_contract.ports["raster"].produces
        == GEO_RASTER_SCAN.key
    )
    assert list(import_raster.node_class.config_contract.model.model_fields) == [
        "source_name"
    ]
    assert import_raster.staged_upload_inputs == ()
    assert import_raster.required_capabilities == (PluginRuntimeCapability.NATIVE_GDAL,)

    assert {spec.key for spec in IMAGES.artifact_type_dependencies} == {
        PNG_FILE.key,
        JPEG_FILE.key,
        WEBP_FILE.key,
        TIFF_FILE.key,
        BMP_FILE.key,
    }
    assert {spec.key for spec in TABLES.artifact_type_dependencies} == {
        CSV_FILE.key,
        XLSX_FILE.key,
    }
    assert {spec.key for spec in GIS.artifact_type_dependencies} == {
        TABLE_DATA.key,
        GEOJSON_FILE.key,
        JSON_FILE.key,
        TIFF_FILE.key,
    }
